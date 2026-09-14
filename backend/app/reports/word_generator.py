from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAVY="173B59"; BLUE="1D5F8F"; PALE="EEF4F8"; GRID="D5DEE6"

def _shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def _set_cell(cell,text,bold=False,color="263B4E",size=9):
    cell.text=""; p=cell.paragraphs[0]; r=p.add_run(str(text if text is not None else "")); r.bold=bold; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP

def _table_borders(table):
    tblPr=table._tbl.tblPr; borders=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        tag=OxmlElement(f"w:{edge}"); tag.set(qn("w:val"),"single"); tag.set(qn("w:sz"),"4"); tag.set(qn("w:color"),GRID); borders.append(tag)
    tblPr.append(borders)

def _app(v): return "APPLICABLE" if v is True else "NOT APPLICABLE"
def _result(t): return str(t.get("result") or "NOT TESTED").upper()

def _heading(doc,text,level=1):
    p=doc.add_paragraph(); r=p.add_run(text); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); r.font.size=Pt(13 if level==1 else 11); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5); return p

def generate_word(report_data):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(9)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("NAWI"); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("NON-AUTOMATIC WEIGHING INSTRUMENT TEST & COMPLIANCE REPORT"); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string(BLUE)
    ev=report_data["evaluation"]; inst=report_data["instrument"]; rng=report_data.get("instrument_range") or {}; tests=report_data.get("tests",[]); summary=report_data.get("summary",{}); attachments=report_data.get("attachments",[]); env=report_data.get("environment")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(f"Overall Result: {summary.get('overall_result','INCOMPLETE')}"); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string({"PASS":"1D7545","FAIL":"B33B31"}.get(summary.get("overall_result"),"9A6A14"))
    doc.add_paragraph(summary.get("overall_statement", ""))
    _heading(doc,"1. Instrument identification")
    table=doc.add_table(rows=0,cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; _table_borders(table)
    for k,v in [("Manufacturer",inst.get("manufacturer")),("Model / Type",inst.get("model") or inst.get("model_type")),("Serial Number",inst.get("serial_number")),("Accuracy Class",inst.get("accuracy_class")),("Instrument Category",inst.get("category")),("Indication Type",inst.get("indication_type")),("Software ID",inst.get("software_identification")),("Software Version",inst.get("software_version"))]:
        c=table.add_row().cells; _set_cell(c[0],k,True,NAVY); _shade(c[0],PALE); _set_cell(c[1],v)
    _heading(doc,"2. Capacity and scale information")
    table=doc.add_table(rows=0,cols=2); _table_borders(table)
    for k,v in [("Max",rng.get("max_capacity")),("Min",rng.get("min_capacity")),("Verification Scale Interval (e)",rng.get("verification_scale_interval")),("Actual Scale Interval (d)",rng.get("actual_scale_interval")),("Maximum Tare",rng.get("maximum_tare"))]:
        c=table.add_row().cells; _set_cell(c[0],k,True,NAVY); _shade(c[0],PALE); _set_cell(c[1],v)
    _heading(doc,"3. Evaluation information")
    for label,value in [("Evaluation Type",ev.get("evaluation_type")),("Status",ev.get("status")),("Evaluation ID",ev.get("id")),("Notes",ev.get("notes") or "—")]: doc.add_paragraph(f"{label}: {value}")
    if env:
        _heading(doc,"4. Test environment")
        for label,value in [("Temperature",env.get("temperature")),("Humidity",env.get("humidity")),("Pressure",env.get("pressure")),("Location",env.get("location")),("Conditions acceptable",env.get("conditions_ok"))]: doc.add_paragraph(f"{label}: {value}")
    _heading(doc,"5. Compliance summary")
    table=doc.add_table(rows=1,cols=2); _table_borders(table); _set_cell(table.rows[0].cells[0],"Metric",True,"FFFFFF"); _set_cell(table.rows[0].cells[1],"Value",True,"FFFFFF"); _shade(table.rows[0].cells[0],NAVY); _shade(table.rows[0].cells[1],NAVY)
    for k,v in [("Applicable tests",summary.get("applicable_tests",0)),("Completed tests",summary.get("completed_tests",0)),("Passed tests",summary.get("passed_tests",0)),("Failed tests",summary.get("failed_tests",0)),("Not tested",summary.get("not_tested_tests",0)),("Not applicable",summary.get("not_applicable_tests",0))]: c=table.add_row().cells; _set_cell(c[0],k); _set_cell(c[1],v)
    _heading(doc,"6. Test result matrix")
    table=doc.add_table(rows=1,cols=5); _table_borders(table)
    for i,h in enumerate(["#","Test","Applicability","Status","Result"]): _set_cell(table.rows[0].cells[i],h,True,"FFFFFF"); _shade(table.rows[0].cells[i],NAVY)
    for i,t in enumerate(tests,1):
        c=table.add_row().cells; r=_result(t); _set_cell(c[0],i); _set_cell(c[1],t.get("test_name") or t.get("test_code")); _set_cell(c[2],_app(t.get("applicability"))); _set_cell(c[3],t.get("status")); _set_cell(c[4],r,True,{"PASS":"1D7545","FAIL":"B33B31"}.get(r,"9A6A14"))
    _heading(doc,"7. Detailed test results")
    for t in tests:
        _heading(doc,str(t.get("test_name") or t.get("test_code") or "OIML Test"),2)
        doc.add_paragraph(f"Applicability: {_app(t.get('applicability'))}")
        doc.add_paragraph(f"Execution Status: {t.get('status','NOT_TESTED')}")
        doc.add_paragraph(f"Result: {_result(t)}")
        if t.get("applicability_reason"): doc.add_paragraph(f"Applicability Reason: {t['applicability_reason']}")
        if t.get("criterion"): doc.add_paragraph(f"Criterion: {t['criterion']}")
        if t.get("explanation"): doc.add_paragraph(f"Explanation: {t['explanation']}")
        if t.get("remarks"): doc.add_paragraph(f"Remarks: {t['remarks']}")
        details=t.get("details") or {}; obs=details.get("observations") if isinstance(details,dict) else None
        if isinstance(obs,list) and obs:
            ot=doc.add_table(rows=1,cols=2); _table_borders(ot); _set_cell(ot.rows[0].cells[0],"#",True,"FFFFFF"); _set_cell(ot.rows[0].cells[1],"Recorded observation",True,"FFFFFF"); _shade(ot.rows[0].cells[0],NAVY); _shade(ot.rows[0].cells[1],NAVY)
            for j,o in enumerate(obs,1):
                c=ot.add_row().cells; _set_cell(c[0],j); text=o if isinstance(o,str) else "\n".join(f"{k}: {v}" for k,v in o.items()); _set_cell(c[1],text)
    if attachments:
        _heading(doc,"8. Evidence and supporting records")
        for a in attachments:
            path=Path(a.get("file_path") or "")
            if a.get("is_image") and a.get("local_exists") and path.exists():
                try:
                    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(5.2)); cap=doc.add_paragraph(f"Figure — {a.get('description') or a.get('file_name')}"); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
                except Exception: doc.add_paragraph(f"Evidence: {a.get('file_name')}")
            else: doc.add_paragraph(f"Supporting file: {a.get('file_name')} — {a.get('description') or 'Supporting document'}")
    _heading(doc,"9. Final statement")
    doc.add_paragraph("This report records the observations, calculations and outcomes stored for the evaluation. Failed tests are retained and explicitly identified; a failed test does not prevent report generation. The overall compliance result is determined from the applicable test outcomes shown in this report.")
    out=BytesIO(); doc.save(out); out.seek(0); return out
